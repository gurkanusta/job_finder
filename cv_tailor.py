#!/usr/bin/env python3
"""
cv_tailor.py — Bir iş ilanına göre master profilden Türkçe .docx CV üretir.

YALAN YOK: Sadece profile/master-profile.md içindeki gerçek bilgileri kullanır.
İlana uyan gerçek beceri/proje/deneyimi öne çıkarır, alakasızı geriye atar.
Yeni teknoloji, deneyim, yıl, sertifika veya seviye UYDURMAZ.

Kullanım:
    # İlan metnini doğrudan ver (en güvenilir):
    python cv_tailor.py --desc-file ilan.txt --company "Makrops" --title "Junior Full Stack Developer"

    # Panodan/metin olarak:
    python cv_tailor.py --desc "buraya ilan metni..." --company "X" --title "Junior Backend"

    # URL'den çekmeyi dene (best-effort; LinkedIn/Kariyer bloklayabilir):
    python cv_tailor.py --url "https://..." --company "X" --title "Junior Backend"

Gerekli:
    pip install anthropic python-docx
    ANTHROPIC_API_KEY ortam değişkeni ayarlı olmalı.
"""

import argparse
import os
import re
import sys
import urllib.request
from pathlib import Path
from typing import List

HERE = Path(__file__).resolve().parent
MASTER_PROFILE = HERE / "profile" / "master-profile.md"
OUTPUT_DIR = HERE / "tailored"
MODEL = os.environ.get("CV_MODEL", "claude-opus-4-8")


# ── İlan metnini elde et ──────────────────────────────────────────────────────
def fetch_job_text(url: str) -> str:
    """URL'den ilan metnini best-effort çeker (HTML etiketlerini temizler)."""
    req = urllib.request.Request(
        url,
        headers={
            "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
            "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/122.0 Safari/537.36",
            "Accept-Language": "tr-TR,tr;q=0.9,en;q=0.8",
        },
    )
    with urllib.request.urlopen(req, timeout=30) as resp:
        html = resp.read().decode("utf-8", "ignore")
    # script/style at, etiketleri temizle
    html = re.sub(r"<(script|style)[^>]*>.*?</\1>", " ", html, flags=re.S | re.I)
    text = re.sub(r"<[^>]+>", " ", html)
    text = re.sub(r"&[a-zA-Z#0-9]+;", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text


# ── Yapılandırılmış CV şeması (Pydantic) ──────────────────────────────────────
def build_models():
    from pydantic import BaseModel

    class Iletisim(BaseModel):
        telefon: str
        email: str
        github: str
        konum: str
        linkedin: str  # yoksa boş string

    class BeceriGrubu(BaseModel):
        kategori: str
        maddeler: List[str]

    class Deneyim(BaseModel):
        sirket: str
        rol: str
        tarih: str
        konum: str
        maddeler: List[str]

    class Proje(BaseModel):
        ad: str
        aciklama: str
        teknolojiler: List[str]

    class Egitim(BaseModel):
        okul: str
        bolum: str
        tarih: str

    class CV(BaseModel):
        ad: str
        unvan: str  # ilana göre uyarlanmış başlık, ör. "Junior Full Stack Developer"
        iletisim: Iletisim
        ozet: str  # ilana göre uyarlanmış 2-4 cümlelik özet
        beceriler: List[BeceriGrubu]
        deneyim: List[Deneyim]
        projeler: List[Proje]
        egitim: List[Egitim]
        diller: List[str]

    return CV


SYSTEM = """Sen, bir iş ilanına göre CV uyarlayan bir asistansın. Katı kurallar:

1. SADECE sana verilen "MASTER PROFİL" içindeki gerçek bilgileri kullan. Kesinlikle
   yeni teknoloji, deneyim, iş, yıl, sertifika, seviye veya beceri UYDURMA.
2. İlana uyan gerçek becerileri, projeleri ve deneyim maddelerini ÖNE ÇIKAR ve
   önce sırala. İlana uymayan gerçek bilgileri geriye at veya kısalt — ama
   silme mecburiyetin yok; asla yalan bir uyum iddia etme.
3. Şirket adları, tarihler, eğitim bilgileri, deneyim süreleri AYNEN korunur.
4. "unvan" alanını ilanın pozisyonuna göre ayarla, ama kişinin gerçekte yapabildiği
   bir şey olmalı (master profildeki becerilere dayanmalı).
5. "ozet" 2-4 cümle, ilana yönelik ama dürüst. Proje açıklamaları master profildeki
   gerçeklere sadık kalmalı, sadece ilana uygun yönleri vurgulanabilir.
6. Dil: Türkçe. Profesyonel, abartısız ton.
7. Telefon, email, github, konum, linkedin alanlarını master profilden aynen al.
   linkedin yoksa boş string bırak."""


def make_cv(master: str, job_text: str, company: str, title: str):
    import anthropic

    CV = build_models()
    client = anthropic.Anthropic()

    user = f"""MASTER PROFİL (tek gerçek kaynak):
---
{master}
---

İŞ İLANI:
Şirket: {company or "(belirtilmedi)"}
Pozisyon: {title or "(belirtilmedi)"}

İlan metni:
---
{job_text or "(ilan metni verilmedi — sadece pozisyon başlığına ve şirkete göre uyarlama yap)"}
---

Yukarıdaki master profili bu ilana göre uyarla ve yapılandırılmış CV üret.
Uydurma yok — sadece master profildeki gerçekleri ilana göre yeniden düzenle."""

    resp = client.messages.parse(
        model=MODEL,
        max_tokens=8000,
        system=SYSTEM,
        messages=[{"role": "user", "content": user}],
        output_format=CV,
    )
    if resp.parsed_output is None:
        raise RuntimeError(f"Model yapılandırılmış çıktı üretemedi. stop_reason={resp.stop_reason}")
    return resp.parsed_output


# ── .docx render ──────────────────────────────────────────────────────────────
def render_docx(cv, out_path: Path):
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    ACCENT = RGBColor(0x1F, 0x3A, 0x5F)

    def heading(text):
        p = doc.add_paragraph()
        r = p.add_run(text.upper())
        r.bold = True
        r.font.size = Pt(12)
        r.font.color.rgb = ACCENT
        p.space_before = Pt(8)
        p.space_after = Pt(2)
        return p

    # Ad + ünvan
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(cv.ad)
    r.bold = True
    r.font.size = Pt(20)
    r.font.color.rgb = ACCENT
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(cv.unvan)
    r2.font.size = Pt(12)

    # İletişim
    i = cv.iletisim
    parcalar = [i.telefon, i.email, i.github, i.konum]
    if i.linkedin.strip():
        parcalar.append(i.linkedin)
    pc = doc.add_paragraph()
    pc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pc.add_run("  |  ".join([x for x in parcalar if x.strip()])).font.size = Pt(9.5)

    # Özet
    heading("Özet")
    doc.add_paragraph(cv.ozet)

    # Beceriler
    heading("Teknik Beceriler")
    for grup in cv.beceriler:
        p = doc.add_paragraph()
        p.add_run(f"{grup.kategori}: ").bold = True
        p.add_run(", ".join(grup.maddeler))

    # Deneyim
    if cv.deneyim:
        heading("İş Deneyimi")
        for d in cv.deneyim:
            p = doc.add_paragraph()
            p.add_run(f"{d.sirket} — {d.rol}").bold = True
            meta = doc.add_paragraph()
            mr = meta.add_run(f"{d.tarih} | {d.konum}")
            mr.italic = True
            mr.font.size = Pt(9.5)
            for m in d.maddeler:
                doc.add_paragraph(m, style="List Bullet")

    # Projeler
    if cv.projeler:
        heading("Projeler")
        for pr in cv.projeler:
            p = doc.add_paragraph()
            p.add_run(f"{pr.ad}").bold = True
            if pr.teknolojiler:
                p.add_run(f"  ({', '.join(pr.teknolojiler)})").italic = True
            doc.add_paragraph(pr.aciklama)

    # Eğitim
    if cv.egitim:
        heading("Eğitim")
        for e in cv.egitim:
            p = doc.add_paragraph()
            p.add_run(f"{e.okul} — {e.bolum}").bold = True
            p.add_run(f"  ({e.tarih})")

    # Diller
    if cv.diller:
        heading("Yabancı Dil")
        doc.add_paragraph(", ".join(cv.diller))

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))


def slug(s: str) -> str:
    s = s.lower()
    tr = str.maketrans("çğıöşü", "cgiosu")
    s = s.translate(tr)
    s = re.sub(r"[^a-z0-9]+", "-", s).strip("-")
    return s or "cv"


def main():
    ap = argparse.ArgumentParser(description="İlana göre Türkçe .docx CV üretir.")
    ap.add_argument("--url", help="İlan URL'i (best-effort çekilir)")
    ap.add_argument("--desc", help="İlan metni (doğrudan)")
    ap.add_argument("--desc-file", help="İlan metnini içeren dosya")
    ap.add_argument("--company", default="", help="Şirket adı")
    ap.add_argument("--title", default="", help="Pozisyon başlığı")
    ap.add_argument("--out", help="Çıktı .docx yolu (opsiyonel)")
    args = ap.parse_args()

    if not MASTER_PROFILE.exists():
        sys.exit(f"HATA: {MASTER_PROFILE} yok. Önce master profili oluştur.")
    master = MASTER_PROFILE.read_text(encoding="utf-8")

    job_text = ""
    if args.desc_file:
        job_text = Path(args.desc_file).read_text(encoding="utf-8")
    elif args.desc:
        job_text = args.desc
    elif args.url:
        try:
            job_text = fetch_job_text(args.url)
            print(f"[i] İlan çekildi ({len(job_text)} karakter).")
        except Exception as e:
            print(f"[uyarı] URL çekilemedi ({e}). Sadece başlık/şirkete göre uyarlanacak.")

    if not os.environ.get("ANTHROPIC_API_KEY"):
        sys.exit("HATA: ANTHROPIC_API_KEY ayarlı değil.")

    print(f"[i] Model: {MODEL} — CV üretiliyor...")
    cv = make_cv(master, job_text, args.company, args.title)

    if args.out:
        out = Path(args.out)
    else:
        name = f"{slug(args.company or 'sirket')}-{slug(args.title or cv.unvan)}.docx"
        out = OUTPUT_DIR / name
    render_docx(cv, out)
    print(f"[✓] CV yazıldı: {out}")


if __name__ == "__main__":
    main()
